#!/usr/bin/env python3
"""Poll the workbook's publishers; write new stories as UTF-8 JSON Lines."""
from __future__ import annotations

import argparse
import asyncio
import hashlib
import json
import logging
import re
import sqlite3
import sys
import time
import unicodedata
from collections import OrderedDict
from contextlib import aclosing
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urljoin, urlsplit, urlunsplit

import aiohttp
from lxml import etree

import extract
from extract import clean, date, document, plain

LOG = logging.getLogger("news")
MAX_BYTES = 5 * 1024 * 1024
USER_AGENT = "TurkishNewsCollector/1.0 (public news research)"
HOST_UPDATES = {"haberglobal.com.tr": "haberglobal.com", "www.haberglobal.com.tr": "haberglobal.com"}
# Click and campaign identifiers that never change the story (newspaper/fundus URL cleaning).
TRACKING = {"fbclid", "gclid", "dclid", "gbraid", "wbraid", "yclid", "msclkid", "mc_cid", "mc_eid", "igshid",
            "_ga", "_gl", "ocid", "cmpid", "xtor", "spm", "twclid", "ttclid"}
# Paths that are never stories: corporate, legal, service and listing pages.
NON_ARTICLE_PATH = re.compile(
    r"/(?:yazarlar|author|etiket|tag|tags|arama|search|kategori|category|iletisim|contact|kunye|hakkimizda|about|"
    r"gizlilik|privacy|cerez|cookie|kvkk|kullanim-kosullari|terms|reklam|advert|abone|subscribe|uyelik|giris|login|"
    r"kayit|register|yayin-akisi|canli-yayin|canli-izle|hava-durumu|namaz-vakitleri|nobetci-eczane|puan-durumu|"
    r"fikstur|burclar|rss|sitemap)(?:/|$)|\.(?:jpg|jpeg|png|gif|webp|pdf|mp4|mp3)$", re.I)


async def bounded_map(function, values, concurrency):
    """Yield completed results with bounded work and cancel pending jobs on exit."""
    values, pending = iter(values), set()
    def fill():
        while len(pending) < concurrency:
            try:
                value = next(values)
            except StopIteration:
                break
            pending.add(asyncio.create_task(function(value)))
    try:
        fill()
        while pending:
            done, pending = await asyncio.wait(pending, return_when=asyncio.FIRST_COMPLETED)
            # Keep all finished tasks tracked until their results are consumed.
            pending.update(done)
            for task in done:
                pending.remove(task)
                yield task.result()
            fill()
    finally:
        for task in pending:
            task.cancel()
        await asyncio.gather(*pending, return_exceptions=True)


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def normalize(value: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", value.casefold().replace("ı", "i"))
                   if not unicodedata.combining(c))


def canonical(url: str, base: str = "") -> str:
    parts = urlsplit(urljoin(base, url))
    if parts.scheme not in {"http", "https"} or not parts.hostname or parts.username:
        raise ValueError("Expected a public HTTP(S) URL")
    if parts.port not in {None, 80, 443}:
        raise ValueError("Unexpected URL port")
    host = parts.hostname.lower()
    if host in {"localhost", "localhost.localdomain"} or ":" in host or re.fullmatch(r"[\d.]+", host):
        raise ValueError("Expected a publisher hostname")
    query = [(k, v) for k, v in parse_qsl(parts.query, keep_blank_values=True)
             if not k.lower().startswith("utm_") and k.lower() not in TRACKING]
    return urlunsplit((parts.scheme, HOST_UPDATES.get(host, parts.netloc.lower()),
                      parts.path or "/", urlencode(query), ""))


def same_site(url: str, base: str) -> bool:
    host = (urlsplit(url).hostname or "").removeprefix("www.")
    domain = (urlsplit(base).hostname or "").removeprefix("www.")
    return host == domain or host.endswith("." + domain)


def load_sources(path: Path) -> list[dict]:
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=False, data_only=True)
        try:
            sheet = next((s for s in book if normalize(s.title) == "kayit defteri"), None)
            if sheet is None:
                raise ValueError('Workbook needs a "Kayıt Defteri" sheet')
            headers = {normalize(str(c.value)): i for i, c in enumerate(sheet[1]) if c.value}
            required = {"kurum", "slug", "adres"}
            if not required <= headers.keys():
                raise ValueError(f"Registry needs these headers: {sorted(required)}")
            sources = []
            for row in sheet.iter_rows(min_row=2):
                address = row[headers["adres"]]
                url = address.hyperlink.target if address.hyperlink else address.value
                if url:
                    sources.append({"name": row[headers["kurum"]].value,
                                    "slug": row[headers["slug"]].value,
                                    "url": url, "feeds": []})
            feeds = next((s for s in book if normalize(s.title) == "canli akislar"), None)
            if feeds:
                columns = {normalize(str(c.value)): i for i, c in enumerate(feeds[1]) if c.value}
                if {"slug", "akis adresi (rss)"} <= columns.keys():
                    by_slug = {s["slug"]: s for s in sources}
                    for row in feeds.iter_rows(min_row=2):
                        source = by_slug.get(row[columns["slug"]].value)
                        cell = row[columns["akis adresi (rss)"]]
                        url = cell.hyperlink.target if cell.hyperlink else cell.value
                        if source and isinstance(url, str) and url.startswith(("http://", "https://")):
                            source["feeds"].append(url)
        finally:
            book.close()
    else:
        sources = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(sources, list) or not sources:
        raise ValueError("Source list must be a nonempty JSON array")
    seen = set()
    for source in sources:
        if not isinstance(source, dict) or not all(isinstance(source.get(k), str) and source[k].strip() for k in ("name", "slug", "url")):
            raise ValueError("Every source needs a name, slug and URL")
        if source["slug"] in seen:
            raise ValueError(f"Duplicate slug: {source['slug']}")
        seen.add(source["slug"])
        source["url"] = canonical(source["url"])
        if not isinstance(source.get("feeds", []), list) or not all(isinstance(u, str) for u in source.get("feeds", [])):
            raise ValueError("Feeds must be an array of URL strings")
        source["feeds"] = list(dict.fromkeys(canonical(u) for u in source.get("feeds", [])))
    return sources


@dataclass
class Article:
    source: str
    source_name: str
    url: str
    title: str
    published_at: str | None
    summary: str = ""
    text: str = ""
    method: str = "rss"
    fetched_at: str = ""
    authors: list[str] = field(default_factory=list)
    keywords: list[str] = field(default_factory=list)
    image: str = ""
    language: str = ""
    section: str = ""
    modified_at: str | None = None
    # Where the body came from: rss, json-ld, publisher, selector, gravity or article.
    content_source: str = ""


def parse_feed(body: bytes, source: dict, base: str) -> list[Article]:
    parser = etree.XMLParser(resolve_entities=False, no_network=True, recover=True)
    root = etree.fromstring(body, parser)
    if root is None or etree.QName(root).localname not in {"rss", "feed", "RDF"}:
        return []
    articles = []
    for item in root.xpath('//*[local-name()="item" or local-name()="entry"]'):
        def field(*names):
            for name in names:
                nodes = item.xpath(f'./*[local-name()="{name}"]')
                if nodes:
                    return "".join(nodes[0].itertext()).strip()
            return ""
        links = item.xpath('./*[local-name()="link"]')
        link = next((n.get("href") or n.text for n in links if n.get("rel", "alternate") == "alternate"), None)
        link = link or field("guid")
        try:
            url = canonical(link or "", base)
        except ValueError:
            continue
        title = plain(field("title"))
        if not link or not title or not same_site(url, source["url"]):
            continue
        content = plain(field("encoded", "content"))
        authors = [clean(n.xpath('string(./*[local-name()="name"])') or n.text) for n in
                   item.xpath('./*[local-name()="creator" or local-name()="author"]')]
        keywords = [clean(n.get("term") or n.text) for n in item.xpath('./*[local-name()="category"]')]
        images = item.xpath('./*[local-name()="content" or local-name()="thumbnail"][@url]/@url | '
                            './*[local-name()="enclosure"][starts-with(@type,"image")]/@url | '
                            './*[local-name()="image"]/text()')
        articles.append(Article(source["slug"], source["name"], url, title,
                                date(field("pubDate", "published", "date", "updated")),
                                plain(field("description", "summary")), content,
                                authors=extract.filter_authors(authors, source["name"]),
                                keywords=list(dict.fromkeys(k for k in keywords if k)),
                                image=next((clean(u) for u in images if clean(u).startswith(("http://", "https://"))), ""),
                                content_source="rss" if content else ""))
    return articles


def discover_feeds(body: bytes, base: str) -> list[str]:
    doc = document(body)
    urls = []
    for node in doc.xpath('//link[@href] | //a[@href]'):
        kind, href = node.get("type", "").lower(), node.get("href", "")
        if "rss" in kind or "atom" in kind or re.search(r"(?:/rss(?:[/.?]|$)|/feed/?$)", href):
            try:
                url = canonical(href, base)
                if same_site(url, base):
                    urls.append(url)
            except ValueError:
                pass
    return list(dict.fromkeys(urls))[:12]


def page_candidates(body: bytes, base: str, limit: int) -> list[str]:
    doc = document(body)
    urls = []
    for node in doc.xpath('//a[@href]'):
        label = clean(node.text_content()) or clean(node.get("title"))
        if not label:
            label = clean(" ".join(node.xpath('.//img/@alt')))
        article_link = bool(node.xpath('.//h1 | .//h2 | .//h3 | .//img'))
        if len(label) < 10 or (len(label) < 25 and not article_link):
            continue
        try:
            url = canonical(node.get("href"), base)
        except ValueError:
            continue
        path = urlsplit(url).path
        story_path = len(path.strip("/")) > 20 or bool(re.search(r"\d{3,}", path))
        if same_site(url, base) and story_path and not NON_ARTICLE_PATH.search(path):
            urls.append(url)
    return list(dict.fromkeys(urls))[:limit]


def parse_article(body: bytes, source: dict, url: str) -> Article | None:
    fields = extract.extract(body, url, source["name"])
    if fields is None:
        return None
    return Article(source["slug"], source["name"], url, method="html", **fields)


def sitemap_links(body: bytes, base: str, limit: int) -> tuple[list[str], list[str]]:
    root = etree.fromstring(body, etree.XMLParser(resolve_entities=False, no_network=True, recover=True))
    if root is None:
        return [], []
    urls = []
    for loc in root.xpath('//*[local-name()="loc"]/text()'):
        try:
            url = canonical(loc, base)
            if same_site(url, base):
                urls.append(url)
        except ValueError:
            pass
    if etree.QName(root).localname == "sitemapindex":
        # Prefer news maps; avoid recursively walking historical archives.
        return [], sorted(urls, key=lambda u: "news" not in u.lower())[:2]
    return urls[:limit], []


class Collector:
    def __init__(self, args):
        self.args = args
        args.db.parent.mkdir(parents=True, exist_ok=True)
        self.db = sqlite3.connect(args.db)
        self.db.execute("PRAGMA journal_mode=WAL")
        self.db.execute('''CREATE TABLE IF NOT EXISTS articles (
            id TEXT PRIMARY KEY, source TEXT, url TEXT, title TEXT, published_at TEXT,
            fetched_at TEXT, payload TEXT NOT NULL, UNIQUE(source, url))''')
        self.cache = OrderedDict()
        self.cache_size = 0
        self.cache_limit = getattr(args, "cache_mb", 32) * 1024 * 1024
        self.endpoint_cooldowns = {}
        self.http_gate = asyncio.Semaphore(args.concurrency)
        self.site_gates = {}
        self.site_slots = {}
        self.site_concurrency = getattr(args, "per_site_concurrency", 2)
        self.last_request = {}

    def seen(self, source: str, url: str) -> bool:
        return self.db.execute("SELECT 1 FROM articles WHERE source=? AND url=?", (source, url)).fetchone() is not None

    def save(self, article: Article) -> bool:
        article.fetched_at = now()
        payload = json.dumps(asdict(article), ensure_ascii=False)
        identifier = hashlib.sha256(f"{article.source}\n{article.url}".encode()).hexdigest()
        with self.db:
            cursor = self.db.execute("INSERT OR IGNORE INTO articles VALUES (?, ?, ?, ?, ?, ?, ?)",
                                     (identifier, article.source, article.url, article.title,
                                      article.published_at, article.fetched_at, payload))
        if cursor.rowcount:
            print(payload, flush=True)
            return True
        return False

    async def fetch(self, session, url, cache=False):
        url = canonical(url)
        host = urlsplit(url).hostname.removeprefix("www.")
        gate = self.site_gates.setdefault(host, asyncio.Lock())
        previous = self.cache.get(url)
        if previous:
            self.cache.move_to_end(url)
        headers = {}
        if cache and previous:
            headers = {key: value for key, value in (("If-None-Match", previous[1]),
                       ("If-Modified-Since", previous[2])) if value}
        slots = self.site_slots.setdefault(host, asyncio.Semaphore(self.site_concurrency))
        for attempt in range(2):
            retry_wait = None
            try:
                async with slots, self.http_gate:
                    async with gate:
                        delay = max(self.args.delay, getattr(self, "request_delays", {}).get(host, 0))
                        await asyncio.sleep(max(0, delay - (time.monotonic() - self.last_request.get(host, 0))))
                        self.last_request[host] = time.monotonic()
                    # Pacing locks protect request starts, not entire response downloads.
                    # Retries release both connection slots before their backoff.
                    async with session.get(url, headers=headers) as response:
                        if response.status == 304 and previous:
                            return previous[0], previous[3]
                        if response.status in {429, 500, 502, 503, 504} and attempt == 0:
                            retry = response.headers.get("Retry-After", "1")
                            retry_wait = min(float(retry) if retry.isdigit() else 1, 15)
                        else:
                            response.raise_for_status()
                            chunks, size = [], 0
                            async for chunk in response.content.iter_chunked(65536):
                                size += len(chunk)
                                if size > getattr(self, "max_bytes", MAX_BYTES):
                                    raise ValueError("Response exceeds configured byte limit")
                                chunks.append(chunk)
                            body, final_url = b"".join(chunks), canonical(str(response.url))
                            if not same_site(final_url, url):
                                raise ValueError(f"Cross-publisher redirect: {final_url}")
                            if cache:
                                old = self.cache.pop(url, None)
                                self.cache_size -= len(old[0]) if old else 0
                                if self.cache_limit and len(body) <= self.cache_limit:
                                    self.cache[url] = (body, response.headers.get("ETag"),
                                                       response.headers.get("Last-Modified"), final_url)
                                    self.cache_size += len(body)
                                while self.cache_size > self.cache_limit:
                                    self.cache_size -= len(self.cache.popitem(last=False)[1][0])
                            return body, final_url
            except (aiohttp.ClientConnectionError, asyncio.TimeoutError):
                if attempt:
                    raise
                retry_wait = 1
            if retry_wait is not None:
                await asyncio.sleep(retry_wait)

    async def collect(self, session, source):
        started = time.monotonic()
        status = {"source": source["slug"], "name": source["name"], "url": source["url"],
                  "checked_at": now(), "status": "no_articles", "found": 0, "new": 0,
                  "methods": [], "errors": []}
        feed_articles = []
        working_feeds = set()
        prefetched_articles = {}
        async def get(url, cache=True):
            if cache and self.endpoint_cooldowns.get(url, 0) > time.monotonic():
                status["errors"].append({"url": url, "error": "Unavailable feed endpoint cooling down for five minutes"})
                return None
            try:
                return await self.fetch(session, url, cache)
            except (aiohttp.ClientError, asyncio.TimeoutError, ValueError, etree.Error) as exc:
                status["errors"].append({"url": url, "error": f"{type(exc).__name__}: {exc}"})
                if cache and isinstance(exc, aiohttp.ClientResponseError) and exc.status in {404, 410}:
                    self.endpoint_cooldowns[url] = time.monotonic() + 300
                return None
        async def feeds(urls):
            items = []
            async def read(url):
                result = await get(url)
                if result:
                    try:
                        parsed = parse_feed(result[0], source, result[1])
                        feed_articles.extend(parsed)
                        if parsed:
                            working_feeds.add(url)
                        return parsed
                    except (etree.Error, ValueError) as exc:
                        status["errors"].append({"url": url, "error": str(exc)})
                return []
            urls = list(dict.fromkeys(urls))
            if not urls:
                return items
            # The first configured feed is normally the publisher's general feed.
            # Avoid downloading every category when it already fills the sample.
            items.extend(await read(urls[0]))
            def enough():
                return not getattr(self.args, "all_feeds", False) and len({a.url for a in items}) >= self.args.limit
            if enough():
                return items
            async with aclosing(bounded_map(read, urls[1:], self.site_concurrency)) as results:
                async for parsed in results:
                    items.extend(parsed)
                    if enough():
                        break
            return items
        async def read_article(url):
            if url in prefetched_articles:
                return prefetched_articles[url]
            result = await get(url, cache=False)
            try:
                return parse_article(result[0], source, url) if result else None
            except (ValueError, TypeError, etree.Error) as exc:
                status["errors"].append({"url": url, "error": f"{type(exc).__name__}: {exc}"})
                return None
        candidate_limit = max(10, self.args.limit * 3)
        async def collect_pages(urls):
            unseen = []
            for url in dict.fromkeys(urls):
                if self.seen(source["slug"], url):
                    status["found"] += 1
                else:
                    unseen.append(url)
            async with aclosing(bounded_map(read_article, unseen, self.site_concurrency)) as results:
                async for article in results:
                    if article:
                        status["found"] += 1
                        status["new"] += self.save(article)
                    if status["new"] >= self.args.limit:
                        break
        async def sitemap_candidates():
            urls = []
            # Google News maps first; Habertürk and NTV paths come from fundus.
            for path in ("/sitemap_news.xml", "/news-sitemap.xml", "/sitemap_google_news.xml",
                         "/sitemaps/news-sitemap.xml", "/sitemap.xml"):
                result = await get(urljoin(source["url"], path))
                if not result:
                    continue
                found, children = sitemap_links(*result, candidate_limit)
                for child in children:
                    sub = await get(child)
                    if sub:
                        found.extend(sitemap_links(*sub, candidate_limit)[0])
                urls.extend(found)
                if urls:
                    break
            return list(dict.fromkeys(urls))[:candidate_limit]
        try:
            async with asyncio.timeout(self.args.site_timeout):
                articles = await feeds(source["feeds"])
                homepage = None
                if not articles:
                    homepage = await get(source["url"])
                    discovered = discover_feeds(*homepage) if homepage else []
                    home_urls = page_candidates(*homepage, candidate_limit) if homepage else []
                    if home_urls and not discovered and not source.get("_html_fallback"):
                        article = await read_article(home_urls[0])
                        prefetched_articles[home_urls[0]] = article
                        if article:
                            source["_html_fallback"] = True
                    guessed = [urljoin(source["url"], path) for path in
                               ("/rss", "/feed/", "/rss.xml", "/rss/anasayfa.xml")]
                    candidates = list(dict.fromkeys(discovered + guessed))
                    # Reuse a working HTML route on later polls rather than guessing
                    # the same dead feed paths every three minutes.
                    if not source.get("_html_fallback"):
                        candidates = [u for u in candidates if u not in source["feeds"]]
                        for offset in range(0, len(candidates), self.site_concurrency):
                            batch = candidates[offset:offset + self.site_concurrency]
                            articles = await feeds(batch)
                            if articles:
                                # Keep only endpoints that actually returned stories.
                                source["feeds"] = [u for u in batch if u in working_feeds]
                                break
                if articles:
                    status["methods"].append("rss")
                    unique = {a.url: a for a in articles}
                    articles = list(unique.values())
                    articles.sort(key=lambda a: a.published_at or "", reverse=True)
                    status["found"] = len(articles)
                    unseen = [a for a in articles[:self.args.limit] if not self.seen(a.source, a.url)]
                    async def enrich(article):
                        if self.args.full_text:
                            enriched = await read_article(article.url)
                            if enriched:
                                if len(enriched.text) > len(article.text):
                                    article.text, article.content_source = enriched.text, enriched.content_source
                                article.summary = article.summary or enriched.summary
                                article.published_at = article.published_at or enriched.published_at
                                article.authors = article.authors or enriched.authors
                                article.keywords = article.keywords or enriched.keywords
                                article.image = article.image or enriched.image
                                for name in ("language", "section", "modified_at"):
                                    setattr(article, name, getattr(enriched, name))
                        return article
                    async with aclosing(bounded_map(enrich, unseen, self.site_concurrency)) as results:
                        async for article in results:
                            status["new"] += self.save(article)
                else:
                    status["methods"].append("html")
                    await collect_pages(home_urls)
                    # Homepage links can lead to categories or incomplete markup.
                    # They must not suppress the sitemap fallback when none parse.
                    if not status["found"]:
                        await collect_pages(await sitemap_candidates())
                    source["_html_fallback"] = bool(status["found"])
                if status["found"]:
                    status["status"] = "ok"
        except TimeoutError:
            if feed_articles and not status["found"]:
                articles = sorted({a.url: a for a in feed_articles}.values(),
                                  key=lambda a: a.published_at or "", reverse=True)
                status["found"] = len(articles)
                status["methods"].append("rss")
                for article in articles[:self.args.limit]:
                    status["new"] += self.save(article)
            status["status"] = "partial" if status["found"] else "timeout"
            status["errors"].append({"error": "Publisher time budget exceeded"})
        except (ValueError, etree.Error, TypeError) as exc:
            status["status"] = "parse_error"
            status["errors"].append({"error": f"{type(exc).__name__}: {exc}"})
        status["seconds"] = round(time.monotonic() - started, 2)
        LOG.info("%-18s %-12s found=%d new=%d", source["slug"], status["status"], status["found"], status["new"])
        return status

    async def run(self, sources):
        timeout = aiohttp.ClientTimeout(total=self.args.timeout)
        connector = aiohttp.TCPConnector(limit=self.args.concurrency, limit_per_host=self.site_concurrency)
        try:
            async with aiohttp.ClientSession(timeout=timeout, connector=connector,
                                            headers={"User-Agent": USER_AGENT}, trust_env=True) as session:
                while True:
                    statuses = await asyncio.gather(*(self.collect(session, s) for s in sources))
                    self.args.status.parent.mkdir(parents=True, exist_ok=True)
                    temporary = self.args.status.with_suffix(".tmp")
                    temporary.write_text(json.dumps({"checked_at": now(), "sources": statuses},
                                                    ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
                    temporary.replace(self.args.status)
                    LOG.info("Round finished: %d/%d sources returned articles; %d new stories",
                             sum(s["status"] == "ok" for s in statuses), len(statuses), sum(s["new"] for s in statuses))
                    if not self.args.interval:
                        return 0 if any(s["found"] for s in statuses) else 2
                    await asyncio.sleep(self.args.interval)
        finally:
            self.db.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sources", type=Path, default=Path(__file__).with_name("sources.json"), help="JSON registry or original XLSX")
    parser.add_argument("--db", type=Path, default=Path("news.sqlite3"))
    parser.add_argument("--status", type=Path, default=Path("source_status.json"))
    parser.add_argument("--interval", type=float, default=0, help="Seconds between rounds; 0 runs once")
    parser.add_argument("--limit", type=int, default=20, help="Maximum items considered per source per round")
    parser.add_argument("--concurrency", type=int, default=16)
    parser.add_argument("--per-site-concurrency", type=int, default=2, help="Concurrent requests per publisher; starts still obey --delay")
    parser.add_argument("--cache-mb", type=int, default=32, help="Maximum in-memory HTTP body cache in MiB (0 disables it)")
    parser.add_argument("--timeout", type=float, default=15, help="Timeout for each HTTP request")
    parser.add_argument("--site-timeout", type=float, default=180, help="Total budget per publisher per round")
    parser.add_argument("--delay", type=float, default=0.25, help="Minimum request spacing per publisher")
    parser.add_argument("--full-text", action="store_true", help="Also visit RSS article pages for body text")
    parser.add_argument("--all-feeds", action="store_true", help="Read every category feed before selecting the latest items")
    parser.add_argument("--only", help="Comma-separated source slugs for a focused run")
    args = parser.parse_args()
    if args.interval < 0 or args.delay < 0 or args.cache_mb < 0 or min(args.limit, args.concurrency, args.per_site_concurrency, args.timeout, args.site_timeout) <= 0:
        parser.error("Limits and timeouts must be positive; interval and delay cannot be negative")
    try:
        sources = load_sources(args.sources)
        if args.only:
            selected = set(args.only.split(","))
            unknown = selected - {s["slug"] for s in sources}
            if unknown:
                parser.error(f"Unknown sources: {', '.join(sorted(unknown))}")
            sources = [s for s in sources if s["slug"] in selected]
    except (OSError, ValueError) as exc:
        parser.error(str(exc))
    logging.basicConfig(level=logging.INFO, format="%(message)s", stream=sys.stderr)
    try:
        return asyncio.run(Collector(args).run(sources))
    except KeyboardInterrupt:
        return 130


if __name__ == "__main__":
    raise SystemExit(main())
